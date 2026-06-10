"""Chat routes with user memory injection."""
from fastapi import APIRouter, HTTPException, Request, UploadFile, File
import uuid
import json
import io
import base64
from typing import Optional
from datetime import datetime, timezone
# Migration: weg von emergentintegrations, hin zu nativer Anthropic-/OpenAI-
# SDK via Compat-Shim. API ist bit-identisch mit der Emergent-Klasse, der
# Shim wählt anhand der ENV-Keys den Provider:
#   ANTHROPIC_API_KEY → Claude (Default ab dieser Migration)
#   OPENAI_API_KEY    → GPT
#   sonst EMERGENT_LLM_KEY → Legacy-Wrapper (übergangsweise)
from lib.llm_provider import LlmChat, UserMessage
import pypdf
import docx

from config import db, EMERGENT_LLM_KEY, logger
from models import ChatMessageIn, ChatSessionCreate
from services import get_current_user, get_user_memory, WLADBOT_SYSTEM_PROMPT
from services_actions import record_user_action
from data import LEADERSHIP_QUOTES

router = APIRouter(prefix="/api", tags=["chat"])

# Max sizes
MAX_PDF_SIZE = 15 * 1024 * 1024   # 15 MB
MAX_DOC_SIZE = 10 * 1024 * 1024   # 10 MB
MAX_TXT_SIZE = 2 * 1024 * 1024    # 2 MB
MAX_IMG_SIZE = 10 * 1024 * 1024   # 10 MB
MAX_EXTRACT_CHARS = 40000


def _extract_pdf_text(data: bytes) -> tuple[str, int, int]:
    reader = pypdf.PdfReader(io.BytesIO(data))
    pages_to_read = min(30, len(reader.pages))
    chunks = []
    for i in range(pages_to_read):
        try:
            chunks.append(reader.pages[i].extract_text() or "")
        except Exception:
            continue
    full_text = "\n\n".join(chunks).strip()
    return full_text, len(reader.pages), pages_to_read


def _extract_docx_text(data: bytes) -> str:
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_txt:
                parts.append(row_txt)
    return "\n".join(parts).strip()


def _truncate(text: str) -> str:
    if len(text) > MAX_EXTRACT_CHARS:
        return text[:MAX_EXTRACT_CHARS] + "\n\n[… gekürzt …]"
    return text


def _handle_pdf(filename: str, data: bytes) -> dict:
    if len(data) > MAX_PDF_SIZE:
        raise HTTPException(status_code=413, detail="PDF zu groß (max 15 MB)")
    text, total_pages, pages_read = _extract_pdf_text(data)
    if not text:
        raise HTTPException(status_code=400, detail="PDF enthält keinen extrahierbaren Text")
    text = _truncate(text)
    return {
        "filename": filename, "type": "pdf",
        "pages": total_pages, "pages_read": pages_read,
        "chars": len(text), "text_preview": text[:400], "full_text": text,
    }


def _handle_docx(filename: str, data: bytes) -> dict:
    if len(data) > MAX_DOC_SIZE:
        raise HTTPException(status_code=413, detail="DOCX zu groß (max 10 MB)")
    text = _extract_docx_text(data)
    if not text:
        raise HTTPException(status_code=400, detail="DOCX enthält keinen Text")
    text = _truncate(text)
    return {
        "filename": filename, "type": "docx",
        "chars": len(text), "text_preview": text[:400], "full_text": text,
    }


def _handle_plaintext(filename: str, ext: str, data: bytes) -> dict:
    if len(data) > MAX_TXT_SIZE:
        raise HTTPException(status_code=413, detail="Textdatei zu groß (max 2 MB)")
    try:
        text = data.decode("utf-8", errors="ignore")
    except Exception:
        raise HTTPException(status_code=400, detail="Datei ist nicht UTF-8 lesbar")
    if not text.strip():
        raise HTTPException(status_code=400, detail="Datei ist leer")
    text = _truncate(text)
    return {
        "filename": filename, "type": ext,
        "chars": len(text), "text_preview": text[:400], "full_text": text,
    }


def _handle_image(filename: str, ext: str, data: bytes) -> dict:
    if len(data) > MAX_IMG_SIZE:
        raise HTTPException(status_code=413, detail="Bild zu groß (max 10 MB)")
    img_b64 = base64.b64encode(data).decode()
    mime = f"image/{'jpeg' if ext == 'jpg' else ext}"
    return {
        "filename": filename, "type": "image",
        "mime": mime, "bytes": len(data),
        "image_b64": img_b64, "data_url": f"data:{mime};base64,{img_b64}",
    }


# Extension → handler dispatcher (keeps the main endpoint a thin router).
_DOC_HANDLERS = {
    "pdf":  lambda fn, ext, data: _handle_pdf(fn, data),
    "docx": lambda fn, ext, data: _handle_docx(fn, data),
    "txt":  _handle_plaintext, "md": _handle_plaintext, "csv": _handle_plaintext,
    "png":  _handle_image, "jpg": _handle_image, "jpeg": _handle_image,
    "webp": _handle_image, "gif": _handle_image,
}


@router.post("/chat/upload-document")
async def chat_upload_document(request: Request, file: UploadFile = File(...)):
    """Universal document upload — accepts PDF, DOCX, TXT, MD, PNG, JPG.
    Returns {filename, type, chars/pages, full_text OR image_b64 for vision}."""
    await get_current_user(request)
    name = (file.filename or "upload").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""

    try:
        data = await file.read()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Datei konnte nicht gelesen werden: {e}")

    handler = _DOC_HANDLERS.get(ext)
    if not handler:
        raise HTTPException(status_code=400, detail=f"Format .{ext} nicht unterstützt. Erlaubt: PDF, DOCX, TXT, MD, PNG, JPG")

    try:
        return handler(file.filename, ext, data)
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Doc upload failed")
        raise HTTPException(status_code=500, detail=f"Upload fehlgeschlagen: {e}")


@router.post("/chat/upload-pdf")
async def chat_upload_pdf(request: Request, file: UploadFile = File(...)):
    """Legacy PDF-only endpoint — proxies to universal /chat/upload-document."""
    return await chat_upload_document(request, file)


@router.post("/chat/sessions")
async def create_chat_session(data: ChatSessionCreate, request: Request):
    user = await get_current_user(request)
    session_id = f"chat_{uuid.uuid4().hex[:12]}"
    doc = {
        "session_id": session_id, "user_id": user["user_id"],
        "title": data.title, "agent": data.agent,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    if data.folder_id:
        # Owner-guard: only stamp folder_id if the user actually owns that folder.
        # Foreign folder_id is silently ignored (no info leak).
        own = await db.folders.find_one(
            {"folder_id": data.folder_id, "user_id": user["user_id"]},
            {"_id": 0, "folder_id": 1},
        )
        if own:
            doc["folder_id"] = data.folder_id
            # Auto-link the new chat session into the folder so it shows up in
            # the folder's item list — mirrors the video-mission behavior.
            await db.folder_items.insert_one({
                "item_id": f"fi_{uuid.uuid4().hex[:12]}",
                "folder_id": data.folder_id,
                "user_id": user["user_id"],
                "item_type": "chat_session",
                "source_id": session_id,
                "title": (data.title or "")[:200],
                "created_at": doc["created_at"],
            })
    await db.chat_sessions.insert_one(doc)
    return {"session_id": session_id, "title": data.title, "agent": data.agent, "folder_id": doc.get("folder_id"), "created_at": doc["created_at"]}


@router.get("/chat/sessions")
async def list_chat_sessions(request: Request, folder_id: Optional[str] = None):
    """List my chat sessions, optionally filtered by folder.

    Iter 92.23.8: when ?folder_id= is set we only return sessions tagged with
    that folder — used by the folder-scoped chat sidebar in the new UX.
    """
    user = await get_current_user(request)
    query: dict = {"user_id": user["user_id"]}
    if folder_id:
        query["folder_id"] = folder_id
    return await db.chat_sessions.find(query, {"_id": 0, "session_id": 1, "title": 1, "agent": 1, "folder_id": 1, "created_at": 1, "updated_at": 1, "message_count": 1}).sort("updated_at", -1).to_list(50)


ROLE_MAP = {
    "Kommunikator": "Du agierst als KOMMUNIKATOR-Rolle. Fokus: Charismatisch auftreten und überzeugen. Nutze die 3 Säulen der Überzeugung (Logos, Ethos, Pathos), den Kommunikationsquadrant und die Feedbackformel.",
    "Manager": "Du agierst als MANAGER-Rolle. Fokus: Effektiver und effizienter arbeiten. Nutze die Entscheidungsmatrix, Delegations-Framework und Priorisierung nach Impact.",
    "Team-Leader": "Du agierst als TEAM-LEADER-Rolle. Fokus: Nachhaltig motivieren und besser delegieren. Nutze Delegation als Befähigung, aktives Zuhören und Motivation-Frameworks.",
    "Psychologe": "Du agierst als PSYCHOLOGE-Rolle. Fokus: Jedem Mitarbeiter individuell und empathisch begegnen. Nutze aktives Zuhören (5 Ebenen), Emotionsregulation und empathische Gesprächsführung.",
    "Problemlöser": "Du agierst als PROBLEMLÖSER-Rolle. Fokus: Konflikte managen und Veränderungen durchsetzen. Nutze die 4 Gesprächstypen, Schwarze Rhetorik Defense und Mediationstechniken.",
}


def _build_system_message(agent: str, user_memory: str, folder_context: str = "") -> str:
    """Build AI system message with agent role, user memory, and folder context."""
    agent_context = ""
    if agent and agent != "auto":
        agent_context = f"\n{ROLE_MAP.get(agent, f'Der User hat die Rolle {agent} gewählt. Fokussiere deine Antwort auf die Spezialität dieser Rolle.')}"
    memory_context = f"\n\n--- USER MEMORY (nutze dies fuer personalisierte Antworten) ---\n{user_memory}\n---" if user_memory else ""
    folder_block = ""
    if folder_context:
        folder_block = (
            "\n\n--- AKTIVER ORDNER-KONTEXT (der User arbeitet gerade in diesem Themenbereich — "
            "alle Antworten sollen diesen Kontext beruecksichtigen) ---\n"
            f"{folder_context}\n---"
        )
    return WLADBOT_SYSTEM_PROMPT + agent_context + memory_context + folder_block


async def _resolve_folder_context(folder_id: Optional[str], user_id: str) -> str:
    """Return the folder's context_summary + recent-item timeline — owner-guarded.

    Iter 92.23.9 (Folder-Knowledge): appends a deterministic timeline of the
    folder's most-recent items so the agent "kennt die Geschichte deines Themas"
    without needing a separate vector store. Foreign folder_id → "" (no leak).
    """
    if not folder_id:
        return ""
    folder = await db.folders.find_one({"folder_id": folder_id, "user_id": user_id}, {"_id": 0})
    if not folder:
        return ""
    summary = (folder.get("context_summary") or "").strip()
    name = folder.get("name") or ""
    parts = [f"Ordner: {name}"]
    if summary:
        parts.append(f"Zusammenfassung: {summary}")
    # Folder-Knowledge: append item timeline (best-effort — never block chat).
    try:
        from services_folder_knowledge import build_folder_timeline
        _items, timeline_text = await build_folder_timeline(folder_id, user_id)
        if timeline_text:
            parts.append(timeline_text)
    except Exception as e:  # noqa: BLE001
        logger.warning("folder_timeline failed: %s", e)
    return "\n\n".join(parts)


async def _ensure_session(session_id: str, user_id: str, message: str, agent: str, folder_id: Optional[str] = None) -> str:
    """Create chat session if needed, return session_id."""
    if session_id:
        return session_id
    new_id = f"chat_{uuid.uuid4().hex[:12]}"
    doc = {
        "session_id": new_id, "user_id": user_id,
        "title": message[:50], "agent": agent,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    if folder_id:
        doc["folder_id"] = folder_id
    await db.chat_sessions.insert_one(doc)
    return new_id


async def _save_ai_tasks(tasks: list, user_id: str, session_id: str):
    """Persist AI-generated tasks to DB."""
    for task in tasks:
        await db.tasks.insert_one({
            "task_id": f"task_{uuid.uuid4().hex[:12]}", "user_id": user_id,
            "title": task.get("title", ""), "description": task.get("description", ""),
            "priority": task.get("priority", "medium"), "status": "pending",
            "source": "ai_coach", "session_id": session_id,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })


@router.post("/chat")
async def send_chat_message(data: ChatMessageIn, request: Request):
    from routes.credits import check_and_deduct_credit
    user = await get_current_user(request)

    credit_result = await check_and_deduct_credit(user, "chat")
    if not credit_result["allowed"]:
        raise HTTPException(status_code=402, detail="no_credits")

    # Owner-guarded folder context — foreign/absent folder_id resolves to "".
    folder_context = await _resolve_folder_context(data.folder_id, user["user_id"])

    session_id = await _ensure_session(data.session_id, user["user_id"], data.message, data.agent, data.folder_id if folder_context else None)

    await db.chat_messages.insert_one({
        "message_id": f"msg_{uuid.uuid4().hex[:12]}", "session_id": session_id,
        "user_id": user["user_id"], "role": "user", "content": data.message,
        "created_at": datetime.now(timezone.utc).isoformat(),
    })

    history = await db.chat_messages.find({"session_id": session_id}, {"_id": 0}).sort("created_at", 1).to_list(20)
    user_memory = await get_user_memory(user["user_id"])
    system_msg = _build_system_message(data.agent, user_memory, folder_context)

    # RAG: retrieve Wlad-specific knowledge chunks based on the user's question.
    # Graceful — if Voyage/Supabase keys are missing or fail, chat continues
    # normally with just the base system prompt.
    from services_rag import retrieve_context
    rag_ctx = await retrieve_context(data.message)
    if rag_ctx["context_block"]:
        system_msg = system_msg + rag_ctx["context_block"]

    try:
        # Shim wählt Provider via ENV — wenn ANTHROPIC_API_KEY gesetzt ist,
        # läuft das auf Claude Haiku 4.5, sonst GPT, sonst Emergent (Legacy).
        chat = LlmChat(api_key=EMERGENT_LLM_KEY, session_id=f"wladbot_{session_id}", system_message=system_msg)
        chat.with_model("openai", "gpt-5.2")  # nur als Hint; Shim override't mit ENV-Preferenz

        context = "".join(f"{'User' if m['role'] == 'user' else 'Assistant'}: {m['content']}\n" for m in history[-10:])
        ai_response = await chat.send_message(UserMessage(text=f"{context}\nUser: {data.message}"))

        try:
            from services_ai_parse import parse_ai_json
            parsed = parse_ai_json(ai_response) or {"insight": ai_response, "strategy": "", "action_steps": [], "simulation_prompt": None, "reflection": "", "tasks": [], "agent_used": data.agent or "General"}
        except Exception:
            parsed = {"insight": ai_response, "strategy": "", "action_steps": [], "simulation_prompt": None, "reflection": "", "tasks": [], "agent_used": data.agent or "General"}

        msg_id = f"msg_{uuid.uuid4().hex[:12]}"
        await db.chat_messages.insert_one({
            "message_id": msg_id, "session_id": session_id, "user_id": user["user_id"],
            "role": "assistant", "content": json.dumps(parsed), "parsed": parsed,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })

        if parsed.get("tasks"):
            await _save_ai_tasks(parsed["tasks"], user["user_id"], session_id)

        await db.chat_sessions.update_one({"session_id": session_id}, {"$set": {"updated_at": datetime.now(timezone.utc).isoformat()}})
        await record_user_action(user["user_id"], "chat_message", metadata={"session_id": session_id, "agent": data.agent})

        return {
            "session_id": session_id,
            "message_id": msg_id,
            "response": parsed,
            "raw": ai_response,
            "rag": {"active": rag_ctx["rag_active"], "chunks": rag_ctx["chunks_count"]},
        }
    except Exception as e:
        logger.error(f"Chat error: {e}")
        raise HTTPException(status_code=500, detail=f"AI service error: {str(e)}")


@router.get("/chat/history/{session_id}")
async def get_chat_history(session_id: str, request: Request):
    user = await get_current_user(request)
    return await db.chat_messages.find({"session_id": session_id, "user_id": user["user_id"]}, {"_id": 0}).sort("created_at", 1).to_list(100)


@router.get("/quote")
async def get_daily_quote():
    import secrets
    return secrets.choice(LEADERSHIP_QUOTES)
